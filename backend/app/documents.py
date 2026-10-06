"""Classify uploaded files for the read_attachment tool.

Three kinds of attachment exist. Images travel to the model as vision blocks and
are never "read". Documents are parsed to text here. Assets (audio, video,
fonts) are binary and have no readable text — they are stored raw so the model
can place them on a webpage with add_webpage_asset. The upload endpoint rejects
anything outside these allowlists before it reaches storage. Document parsing is
best-effort: a file that looks right but cannot be read raises
UnsupportedDocument with a message the API can return.
"""

from __future__ import annotations

import io
import mimetypes
from pathlib import Path

# Images travel to the model as vision blocks; they are never "read".
IMAGE_MIMES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}

_DOCUMENT_MIMES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Code and plain-text formats read verbatim.
_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".rst",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".html", ".htm", ".css", ".xml",
    ".yaml", ".yml", ".ini", ".cfg", ".toml", ".sql", ".sh", ".java",
    ".c", ".h", ".cpp", ".go", ".rs", ".rb", ".php", ".tex",
}

DOCUMENT_EXTENSIONS = set(_DOCUMENT_MIMES) | _TEXT_EXTENSIONS

# Binary files a webpage can use (audio/video players, embedded fonts, favicon).
# They hold no text to read; the model copies them into the site folder with
# add_webpage_asset. Kept next to the other allowlists so the upload endpoint,
# the picker and the tools agree on what counts as an asset.
ASSET_MIMES = {
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".ogg": "audio/ogg",
    ".m4a": "audio/mp4",
    ".aac": "audio/aac",
    ".flac": "audio/flac",
    ".mp4": "video/mp4",
    ".webm": "video/webm",
    ".mov": "video/quicktime",
    ".woff": "font/woff",
    ".woff2": "font/woff2",
    ".ttf": "font/ttf",
    ".otf": "font/otf",
    ".ico": "image/x-icon",
}

# Cap the stored text so an enormous document can't fill the DB or the model's
# context. read_attachment pages through whatever is under this limit.
MAX_EXTRACTED_CHARS = 200_000


class UnsupportedDocument(ValueError):
    """The file isn't a format we can turn into text."""


def extension_of(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def is_image(filename: str) -> bool:
    return extension_of(filename) in IMAGE_MIMES


def is_document(filename: str) -> bool:
    return extension_of(filename) in DOCUMENT_EXTENSIONS


def is_asset(filename: str) -> bool:
    return extension_of(filename) in ASSET_MIMES


def human_size(num_bytes: int) -> str:
    """A short, human-readable size like "3.2 MB" for pointers and tool output."""
    if num_bytes < 1024:
        return f"{num_bytes} B"
    if num_bytes < 1024 * 1024:
        return f"{round(num_bytes / 1024)} KB"
    return f"{num_bytes / (1024 * 1024):.1f} MB"


def mime_for(filename: str) -> str:
    ext = extension_of(filename)
    if ext in IMAGE_MIMES:
        return IMAGE_MIMES[ext]
    if ext in _DOCUMENT_MIMES:
        return _DOCUMENT_MIMES[ext]
    if ext in ASSET_MIMES:
        return ASSET_MIMES[ext]
    guessed, _ = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


# Vision models see at most ~2 megapixels of detail, so a larger upload only
# costs bandwidth, disk and prompt tokens. Anything above this is downscaled on
# the way in; small images pass through untouched.
MAX_IMAGE_PIXELS = 2_000_000


def _encode_image(image, mime: str) -> bytes:
    """Re-encode a Pillow image to ``mime``, preserving alpha where it exists."""
    buffer = io.BytesIO()
    if mime == "image/jpeg":
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")
        image.save(buffer, format="JPEG", quality=85, optimize=True)
    elif mime == "image/webp":
        image.save(buffer, format="WEBP", quality=85)
    else:
        image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def downscale_image(
    data: bytes, mime: str, max_pixels: int = MAX_IMAGE_PIXELS
) -> bytes:
    """Shrink an oversized image to at most ``max_pixels``, keeping its format.

    Best-effort: the original bytes come back unchanged when the image already
    fits, is animated, or cannot be decoded (so an odd-but-valid upload is never
    rejected here).
    """
    try:
        from PIL import Image, ImageOps
    except ImportError:  # defensive: Pillow is in requirements.txt
        return data
    try:
        with Image.open(io.BytesIO(data)) as image:
            if getattr(image, "is_animated", False):
                return data
            image = ImageOps.exif_transpose(image)
            width, height = image.size
            if width * height <= max_pixels:
                return data
            scale = (max_pixels / (width * height)) ** 0.5
            resized = image.resize(
                (max(1, round(width * scale)), max(1, round(height * scale))),
                Image.Resampling.LANCZOS,
            )
            return _encode_image(resized, mime)
    except Exception:
        return data


def _truncate(text: str) -> str:
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    return (
        text[:MAX_EXTRACTED_CHARS]
        + f"\n\n[Truncated at {MAX_EXTRACTED_CHARS} characters.]"
    )


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p for p in pages if p.strip())


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        parts.append(f"--- Slide {index} ---")
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append("\t".join(cells))
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        parts.append(text)
        notes = getattr(slide, "notes_slide", None)
        if notes is not None:
            note_text = notes.notes_text_frame.text
            if note_text.strip():
                parts.append(f"[Notes] {note_text.strip()}")
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    parts.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(parts)


def _plain_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """Return the document's text, capped at MAX_EXTRACTED_CHARS.

    Raises UnsupportedDocument when the extension isn't handled or the bytes
    can't be parsed as that format.
    """
    ext = extension_of(filename)
    if ext not in DOCUMENT_EXTENSIONS:
        raise UnsupportedDocument("That file type isn't supported.")
    try:
        if ext == ".pdf":
            text = _pdf_text(data)
        elif ext == ".docx":
            text = _docx_text(data)
        elif ext == ".pptx":
            text = _pptx_text(data)
        elif ext == ".xlsx":
            text = _xlsx_text(data)
        else:
            text = _plain_text(data)
    except Exception as exc:  # malformed/encrypted file, parser-specific errors
        raise UnsupportedDocument(f"Couldn't read that {ext} file ({exc}).") from exc
    return _truncate(text.strip())
