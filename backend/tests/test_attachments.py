import io
import json

from app import db, documents

FAKE_PNG = b"\x89PNG\r\n\x1a\nfake-image-bytes"


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes(text: str) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes(text: str) -> bytes:
    """Build a minimal one-page PDF whose content stream draws `text`."""
    content = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
        b" /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(out)


# ---------- extraction ----------

def test_extract_plain_text():
    assert documents.extract_text("notes.txt", b"hello world") == "hello world"


def test_extract_pdf():
    text = documents.extract_text("report.pdf", _pdf_bytes("Hello PDF"))
    assert "Hello PDF" in text


def test_extract_docx():
    text = documents.extract_text("essay.docx", _docx_bytes("Hello Word"))
    assert "Hello Word" in text


def test_extract_pptx():
    text = documents.extract_text("deck.pptx", _pptx_bytes("Hello Slides"))
    assert "Hello Slides" in text


def test_extract_xlsx():
    text = documents.extract_text("grades.xlsx", _xlsx_bytes([["Name", "Score"], ["Ada", "10"]]))
    assert "Ada" in text and "10" in text


def test_extract_rejects_unknown_and_malformed():
    import pytest

    with pytest.raises(documents.UnsupportedDocument):
        documents.extract_text("archive.zip", b"PK\x03\x04")
    with pytest.raises(documents.UnsupportedDocument):
        documents.extract_text("broken.docx", b"not a real docx")


def test_extract_truncates_oversized_text(monkeypatch):
    monkeypatch.setattr(documents, "MAX_EXTRACTED_CHARS", 10)
    text = documents.extract_text("big.txt", b"a" * 50)
    assert text.startswith("a" * 10)
    assert "Truncated" in text


def test_is_image_and_is_document():
    assert documents.is_image("photo.JPG")
    assert not documents.is_document("photo.jpg")
    assert documents.is_document("paper.PDF")
    assert documents.is_document("code.py")
    assert not documents.is_document("song.mp3")


# ---------- upload endpoint ----------

def _chat(signed_in):
    return signed_in.post("/api/v1/chats", json={"mode": "chat"}).json()


def test_upload_document_extracts_text(signed_in):
    chat = _chat(signed_in)
    res = signed_in.post(
        "/api/v1/attachments",
        files={"file": ("notes.txt", b"the answer is 42", "text/plain")},
        data={"chatId": chat["id"]},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["kind"] == "document"
    assert body["filename"] == "notes.txt"
    assert body["charCount"] == len("the answer is 42")

    user = db.get_user_by_email("kid@example.com")
    row = db.get_attachment(user["id"], body["id"])
    assert row["text"] == "the answer is 42"
    assert db.attachment_file_path(row["id"], row["filename"]).read_bytes() == b"the answer is 42"


def test_upload_image_keeps_bytes_no_text(signed_in):
    chat = _chat(signed_in)
    res = signed_in.post(
        "/api/v1/attachments",
        files={"file": ("picture.png", FAKE_PNG, "image/png")},
        data={"chatId": chat["id"]},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["kind"] == "image"
    assert body["charCount"] == 0


def test_upload_rejects_unsupported_type(signed_in):
    chat = _chat(signed_in)
    res = signed_in.post(
        "/api/v1/attachments",
        files={"file": ("clip.mp3", b"\x00\x01", "audio/mpeg")},
        data={"chatId": chat["id"]},
    )
    assert res.status_code == 415


def test_upload_rejects_unknown_chat(signed_in):
    res = signed_in.post(
        "/api/v1/attachments",
        files={"file": ("notes.txt", b"hi", "text/plain")},
        data={"chatId": "nope"},
    )
    assert res.status_code == 404


def test_upload_requires_login(signed_in):
    chat = _chat(signed_in)
    signed_in.post("/api/v1/auth/logout")
    res = signed_in.post(
        "/api/v1/attachments",
        files={"file": ("notes.txt", b"hi", "text/plain")},
        data={"chatId": chat["id"]},
    )
    assert res.status_code == 401


# ---------- serving / deletion ----------

def test_serve_attachment_is_user_scoped(signed_in):
    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "notes.txt", b"secret", kind="document",
        mime="text/plain", text="secret",
    )

    served = signed_in.get(f"/api/v1/attachments/{row['id']}")
    assert served.status_code == 200
    assert served.content == b"secret"

    stranger = db.create_google_user("stranger@example.com", "Stranger", "sub-at")
    stranger_chat = db.create_chat(stranger["id"], title="Theirs")
    theirs = db.create_attachment(
        stranger["id"], stranger_chat["id"], "theirs.txt", b"nope",
        kind="document", mime="text/plain", text="nope",
    )
    assert signed_in.get(f"/api/v1/attachments/{theirs['id']}").status_code == 404


def test_delete_attachment_endpoint(signed_in):
    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "notes.txt", b"bye", kind="document",
        mime="text/plain", text="bye",
    )
    path = db.attachment_file_path(row["id"], row["filename"])
    assert path.is_file()

    assert signed_in.delete(f"/api/v1/attachments/{row['id']}").json() == {"ok": True}
    assert db.get_attachment(user["id"], row["id"]) is None
    assert not path.exists()


def test_chat_delete_removes_attachment_rows_and_files(signed_in):
    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "notes.txt", b"bye", kind="document",
        mime="text/plain", text="bye",
    )
    path = db.attachment_file_path(row["id"], row["filename"])
    assert path.is_file()

    assert signed_in.delete(f"/api/v1/chats/{chat['id']}").json() == {"ok": True}
    assert db.get_attachment(user["id"], row["id"]) is None
    assert not path.exists()


# ---------- read_attachment tool ----------

def test_read_attachment_tool_returns_text(signed_in):
    from agent import tools

    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "notes.txt", b"the answer is 42",
        kind="document", mime="text/plain", text="the answer is 42",
    )
    out = tools.read_attachment.invoke(
        {"attachment_id": row["id"]},
        config={"configurable": {"user_id": user["id"]}},
    )
    assert "notes.txt" in out
    assert "the answer is 42" in out


def test_read_attachment_tool_pages_and_scopes(signed_in, monkeypatch):
    from agent import tools

    monkeypatch.setattr(tools, "ATTACHMENT_CHUNK", 5)
    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "big.txt", b"0123456789",
        kind="document", mime="text/plain", text="0123456789",
    )
    first = tools.read_attachment.invoke(
        {"attachment_id": row["id"]},
        config={"configurable": {"user_id": user["id"]}},
    )
    assert "01234" in first and "offset=5" in first
    rest = tools.read_attachment.invoke(
        {"attachment_id": row["id"], "offset": 5},
        config={"configurable": {"user_id": user["id"]}},
    )
    assert "56789" in rest

    other = tools.read_attachment.invoke(
        {"attachment_id": row["id"]},
        config={"configurable": {"user_id": "someone-else"}},
    )
    assert "couldn't find" in other.lower()


def test_read_attachment_tool_rejects_image(signed_in):
    from agent import tools

    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "pic.png", FAKE_PNG, kind="image",
        mime="image/png",
    )
    out = tools.read_attachment.invoke(
        {"attachment_id": row["id"]},
        config={"configurable": {"user_id": user["id"]}},
    )
    assert "image" in out.lower()


# ---------- prompt-state expansion ----------

def _chat_request(chat_id, parts):
    import app.main as main

    return main.ChatRequest(
        chatId=chat_id,
        messages=[{"role": "user", "content": "look at this", "parts": parts}],
    )


def test_state_puts_image_attachment_in_vision_block(signed_in):
    import app.main as main

    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "pic.png", FAKE_PNG, kind="image",
        mime="image/png",
    )
    req = _chat_request(
        chat["id"],
        [{"type": "attachment", "attachmentId": row["id"], "kind": "image"}],
    )
    state = main._state(req, user)
    content = state["messages"][0].content
    assert content[0] == {"type": "text", "text": "look at this"}
    image_block = content[1]
    assert image_block["type"] == "image_url"
    assert image_block["image_url"]["url"].startswith("data:image/png;base64,")


def test_state_points_to_document_for_tool(signed_in):
    import app.main as main

    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "notes.txt", b"hi there", kind="document",
        mime="text/plain", text="hi there",
    )
    req = _chat_request(
        chat["id"],
        [{"type": "attachment", "attachmentId": row["id"], "kind": "document"}],
    )
    state = main._state(req, user)
    content = state["messages"][0].content
    pointer = content[1]["text"]
    assert "notes.txt" in pointer
    assert row["id"] in pointer


def test_state_without_attachments_is_plain_text(signed_in):
    import app.main as main

    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    req = main.ChatRequest(
        chatId=chat["id"], messages=[{"role": "user", "content": "just text"}]
    )
    state = main._state(req, user)
    assert state["messages"][0].content == "just text"


def test_stream_passes_attachment_state_to_graph(signed_in, monkeypatch):
    import app.main as main
    from types import SimpleNamespace

    chat = _chat(signed_in)
    user = db.get_user_by_email("kid@example.com")
    row = db.create_attachment(
        user["id"], chat["id"], "pic.png", FAKE_PNG, kind="image",
        mime="image/png",
    )
    seen = {}

    class FakeGraph:
        async def astream_events(self, state, config=None, version=None):
            seen["messages"] = state["messages"]
            yield {
                "event": "on_chat_model_stream",
                "data": {"chunk": SimpleNamespace(content="I see it!")},
            }

    monkeypatch.setattr(main, "GRAPH", FakeGraph())
    res = signed_in.post(
        "/api/v1/chat/stream",
        json={
            "chatId": chat["id"],
            "messages": [
                {
                    "role": "user",
                    "content": "look",
                    "parts": [
                        {
                            "type": "attachment",
                            "attachmentId": row["id"],
                            "kind": "image",
                        }
                    ],
                }
            ],
            "generateTitle": False,
        },
    )
    assert res.status_code == 200
    content = seen["messages"][0].content
    assert any(block.get("type") == "image_url" for block in content)
